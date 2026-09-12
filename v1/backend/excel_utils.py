"""
excel_utils.py — Utilitaires pour la manipulation de fichiers Excel.
"""
import os
from openpyxl import Workbook, load_workbook


def get_or_create_empty_sheet(filepath: str):
    """
    Cherche la première feuille vide d'un classeur Excel, ou en crée une nouvelle.
    Si le fichier n'existe pas, il est créé.

    Retourne:
        tuple: (Workbook, Worksheet) pour permettre la sauvegarde après modification.
    """
    if os.path.exists(filepath):
        wb = load_workbook(filepath)
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            # Dans openpyxl, une feuille vierge a max_row=1, max_column=1 et A1=None
            if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None:
                return wb, ws

        # Si aucune feuille vide n'est trouvée, on en ajoute une nouvelle
        ws = wb.create_sheet()
        #len(wb.sheetnames)
        return wb, ws
    else:
        # Le fichier n'existe pas, on le crée avec sa feuille par défaut
        wb = Workbook()
        ws = wb.active
        return wb, ws